"""Exercise filing-builder behavior using invented parties and temporary outputs.

These checks inspect DOCX content/structure, not rendered layout or legal validity.
"""

import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document


SCRIPT = (Path(__file__).resolve().parents[1] / "skills" /
          "wisconsin-legal-writing" / "scripts" / "build_filing.py")
MODULE_SPEC = importlib.util.spec_from_file_location("filing_builder_test", SCRIPT)
builder = importlib.util.module_from_spec(MODULE_SPEC)
MODULE_SPEC.loader.exec_module(builder)


def filing():
    return {
        "county": "Synthetic",
        "case_no": "TEST-1",
        "title": "Synthetic motion for testing",
        "first_party": {"name": "Alex Example", "designation": "Plaintiff"},
        "second_party": {"name": "Taylor Example", "designation": "Defendant"},
        "body": [{"type": "p", "text": "This is a synthetic request, not legal advice."}],
        "signature": {
            "name": "Alex Example", "role": "Plaintiff, pro se",
            "address": "1 Synthetic Way", "phone": "555-0100",
            "email": "alex@example.invalid", "date": "January 2, 2026",
        },
    }


class FilingBuilderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.output = self.root / "review.docx"

    def cli(self, spec, *options):
        source = self.root / "spec.json"
        source.write_text(json.dumps(spec), encoding="utf-8-sig")
        return subprocess.run(
            [sys.executable, str(SCRIPT), str(source), "-o", str(self.output), *options],
            capture_output=True, text=True, encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )

    def text(self):
        return "\n".join(p.text for p in Document(self.output).paragraphs)

    def test_complete_spec_builds_without_draft_and_preserves_supplied_text(self):
        spec = filing()
        self.assertEqual(builder.build(spec, self.output), self.output)
        text = self.text()
        self.assertIn(spec["body"][0]["text"], text)
        self.assertIn("ALEX EXAMPLE,", text)
        self.assertIn("Case No. TEST-1", text)
        self.assertIn("Electronically signed by Alex Example", text)
        self.assertIn("alex@example.invalid", text)
        section = Document(self.output).sections[0]
        self.assertAlmostEqual(section.page_width.inches, 8.5)
        self.assertAlmostEqual(section.page_height.inches, 11)

    def test_cli_accepts_utf8_bom_and_discloses_remaining_review(self):
        result = self.cli(filing())
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertTrue(self.output.is_file())
        self.assertIn("review", result.stderr.lower())
        self.assertIn("layout", result.stderr.lower())

    def test_existing_output_is_preserved_until_overwrite_is_explicit(self):
        original = b"Existing user document must survive."
        self.output.write_bytes(original)
        with self.assertRaises(FileExistsError):
            builder.build(filing(), self.output)
        self.assertEqual(self.output.read_bytes(), original)
        builder.build(filing(), self.output, overwrite=True)
        self.assertIn("synthetic request", self.text())

    def test_invalid_input_preserves_existing_output_even_with_overwrite(self):
        original = b"Existing output."
        self.output.write_bytes(original)
        for changes in (
            {"body": []}, {"body": [{"type": "unknown", "text": "Test"}]},
            {"signature": None}, {"appellate": True}, {"line_spacing": float("nan")},
            {"line_spacing": True}, {"font": "Imaginary Font"},
        ):
            with self.subTest(changes=changes):
                spec = filing()
                spec.update(changes)
                with self.assertRaises(ValueError):
                    builder.build(spec, self.output, overwrite=True)
                self.assertEqual(self.output.read_bytes(), original)

    def test_placeholders_need_draft_and_remain_visible_in_draft(self):
        spec = filing()
        spec["body"][0]["text"] = "On [DATE], see (Doc. __:__)."
        spec["signature"].pop("phone")
        with self.assertRaises(ValueError):
            builder.build(spec, self.output)
        self.assertFalse(self.output.exists())
        builder.build(spec, self.output, draft=True)
        self.assertIn("On [DATE], see (Doc. __:__).", self.text())

    def test_generated_non_draft_does_not_invent_blank_attorney_date(self):
        spec = filing()
        spec["signature"] = {
            "style": "attorney", "firm": "Synthetic Firm", "name": "Alex Example",
            "bar_no": "0000000", "for_party": "Attorney for synthetic plaintiff",
            "address": "1 Synthetic Way", "phone": "555-0100", "email": "test@example.invalid",
        }
        try:
            builder.build(spec, self.output)
        except ValueError as exc:
            # Requiring the missing date is also a truthful, supported outcome.
            self.assertIn("date", str(exc).lower())
        else:
            self.assertNotIn("____", self.text())

    def test_malformed_scalar_fields_are_cli_errors_without_tracebacks(self):
        for changes in (
            {"county": 7}, {"title": 7}, {"case_no": 7},
            {"first_party": {"name": 7, "designation": "Plaintiff"}},
        ):
            with self.subTest(changes=changes):
                spec = filing()
                spec.update(changes)
                result = self.cli(spec, "--draft")
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertFalse(self.output.exists())

    def test_proposed_order_allows_no_signature_and_does_not_add_one(self):
        spec = filing()
        spec["proposed_order"] = True
        spec["signature"] = None
        spec["body"] = [{"type": "ordered", "text": "Synthetic ordered paragraph."}]
        builder.build(spec, self.output)
        self.assertIn("Synthetic ordered paragraph.", self.text())
        self.assertNotIn("Electronically signed", self.text())

    def test_legal_omissions_and_literal_asterisks_survive_emphasis(self):
        spec = filing()
        spec["body"][0]["text"] = r"*Synthetic Case*; **emphasis**; 5 * 3; * * *; \*literal\*; __."
        builder.build(spec, self.output, draft=True)
        body = next(p for p in Document(self.output).paragraphs if "Synthetic Case" in p.text)
        self.assertEqual(body.text, "Synthetic Case; emphasis; 5 * 3; * * *; *literal*; __.")
        self.assertTrue(any(run.italic and run.text == "Synthetic Case" for run in body.runs))
        self.assertTrue(any(run.bold and run.text == "emphasis" for run in body.runs))

    def test_non_docx_output_is_refused(self):
        with self.assertRaisesRegex(ValueError, "docx"):
            builder.build(filing(), self.root / "review.pdf")
        self.assertFalse((self.root / "review.pdf").exists())

    def test_output_symlink_never_overwrites_its_target(self):
        target = self.root / "existing.docx"
        target.write_bytes(b"Original linked document")
        try:
            self.output.symlink_to(target)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"symlink creation unavailable: {exc}")
        with self.assertRaises(ValueError):
            builder.build(filing(), self.output, overwrite=True, draft=True)
        self.assertEqual(target.read_bytes(), b"Original linked document")


if __name__ == "__main__":
    unittest.main()
