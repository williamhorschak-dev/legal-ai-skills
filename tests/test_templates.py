"""Build every supplied JSON template; structure and text are not rendered pages."""
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document

from test_builder import builder, filing

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "skills/wisconsin-legal-writing/assets"


class TemplateTests(unittest.TestCase):
    def test_all_json_templates_build_and_preserve_reviewed_document_text(self):
        specs = sorted(ASSETS.glob("*.json"))
        self.assertEqual(len(specs), 11)
        with tempfile.TemporaryDirectory() as directory:
            for source in specs:
                with self.subTest(template=source.stem):
                    spec = json.loads(source.read_text(encoding="utf-8"))
                    target = Path(directory) / f"{source.stem}.docx"
                    builder.build(spec, target, draft=True)
                    document = Document(target)
                    bundled = Document(source.with_suffix(".docx"))
                    self.assertEqual([p.text for p in document.paragraphs], [p.text for p in bundled.paragraphs])
                    section = document.sections[0]
                    self.assertAlmostEqual(section.page_width.inches, 8.5)
                    self.assertAlmostEqual(section.page_height.inches, 11)
                    self.assertAlmostEqual(section.left_margin.inches, 1.0)
                    # Page-one clearance is provided by a distinct first-page
                    # header, not a larger top margin on every page.
                    reviewed_section = bundled.sections[0]
                    self.assertEqual(section.top_margin, reviewed_section.top_margin)
                    self.assertEqual(section.bottom_margin, reviewed_section.bottom_margin)
                    self.assertEqual(section.different_first_page_header_footer, reviewed_section.different_first_page_header_footer)
                    self.assertEqual(section.first_page_header._element.xml, reviewed_section.first_page_header._element.xml)
                    with zipfile.ZipFile(target) as archive:
                        self.assertIsNone(archive.testzip())

    def test_all_word_assets_have_a_content_bound_review_record(self):
        record = json.loads((ROOT / "docs/template-review.json").read_text(encoding="utf-8"))
        assets = {path.name: path for path in ASSETS.glob("*.docx")}
        self.assertEqual(len(assets), 14)
        self.assertEqual(set(assets), set(record["files"]))
        for name, path in assets.items():
            with self.subTest(template=name):
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), record["files"][name]["sha256"],
                                 "Asset changed: render and review it, then update its review record.")
                self.assertTrue(record["files"][name]["review_basis"])
                self.assertTrue(Document(path).paragraphs)

    def test_long_unicode_names_and_signature_remain_intact(self):
        spec = filing()
        spec["first_party"]["name"] = "Renée Example " * 14
        spec["signature"]["name"] = "Renée Example — Jordan Example"
        spec["body"][0]["text"] = "Synthetic Unicode: café, naïve, §, ¶, and an em dash — retained."
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "stress.docx"
            builder.build(spec, target)
            text = "\n".join(p.text for p in Document(target).paragraphs)
            self.assertIn(spec["first_party"]["name"].strip().upper(), text)
            self.assertIn(spec["signature"]["name"], text)
            self.assertIn(spec["body"][0]["text"], text)


if __name__ == "__main__":
    unittest.main()
